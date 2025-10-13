import os, io, re, json
from typing import Tuple, Optional
from bs4 import BeautifulSoup
from pdfminer.high_level import extract_text as pdf_extract
from pypdf import PdfReader
from PIL import Image
from .utils import win_long
import pytesseract

PREFERRED_ENCODINGS = ["utf-8", "gb18030", "gbk", "big5", "cp936", "latin-1"]

def read_text_auto(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext in [".md", ".markdown", ".txt", ".rst", ".tex", ".csv", ".tsv", ".json", ".yaml", ".yml", ".ini", ".log", ".xml"]:
        return read_text_fallback(path)
    elif ext in [".html", ".htm"]:
        return read_html(path)
    elif ext in [".pdf"]:
        return read_pdf(path)
    elif ext in [".docx"]:
        return read_docx(path)
    elif ext in [".xlsx", ".xls"]:
        return read_xlsx(path)
    elif ext in [".pptx"]:
        return read_pptx(path)
    else:
        return ""

def read_text_fallback(path: str) -> str:
    for enc in PREFERRED_ENCODINGS:
        try:
            with open(win_long(path), "r", encoding=enc, errors="strict") as f:
                return f.read()
        except Exception:
            continue
    # 最后兜底忽略错误
    with open(win_long(path), "r", encoding="utf-8", errors="ignore") as f:
        return f.read()

def read_html(path: str) -> str:
    raw = read_text_fallback(path)
    soup = BeautifulSoup(raw, "lxml")
    # 去脚本/样式
    for s in soup(["script","style","noscript"]):
        s.extract()
    # 提取标题与正文
    text = []
    for h in soup.find_all(re.compile("^h[1-6]$")):
        text.append("#" * int(h.name[1]) + " " + h.get_text(" ", strip=True))
    text.append(soup.get_text(" ", strip=True))
    return "\n".join(text)

def read_pdf(path: str) -> str:
    try:
        return pdf_extract(win_long(path)) or ""
    except Exception:
        # 备用库
        try:
            reader = PdfReader(win_long(path))
            out = []
            for page in reader.pages:
                out.append(page.extract_text() or "")
            return "\n".join(out)
        except Exception:
            return ""

def read_docx(path: str) -> str:
    try:
        from docx import Document
        doc = Document(win_long(path))
        return "\n".join(p.text for p in doc.paragraphs)
    except Exception:
        return ""

def read_xlsx(path: str) -> str:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(win_long(path), read_only=True, data_only=True)
        texts = []
        for ws in wb.worksheets:
            texts.append(f"# {ws.title}")
            for row in ws.iter_rows(values_only=True):
                row_txt = "\t".join("" if v is None else str(v) for v in row)
                texts.append(row_txt)
        return "\n".join(texts)
    except Exception:
        return ""

def read_pptx(path: str) -> str:
    try:
        from pptx import Presentation
        prs = Presentation(win_long(path))
        texts = []
        for i, slide in enumerate(prs.slides):
            texts.append(f"# Slide {i+1}")
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    texts.append(shape.text)
        return "\n".join(texts)
    except Exception:
        return ""

def ocr_image(path: str, lang: str="chi_sim+eng", psm: int=6) -> str:
    try:
        img = Image.open(win_long(path))
        cfg = f"--psm {psm}"
        return pytesseract.image_to_string(img, lang=lang, config=cfg) or ""
    except Exception:
        return ""
